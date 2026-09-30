import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = RGBColor(27, 54, 93) # Deep Navy
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(41, 105, 176) # Accent Blue
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(60, 60, 60)
    return p

def add_body_paragraph(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(40, 40, 40)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(50, 50, 50)
    return p

def add_bullet_point(doc, text, bold_prefix=None, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    
    # Bullet symbol
    r_bullet = p.add_run("•  ")
    r_bullet.font.name = "Calibri"
    r_bullet.font.size = Pt(10)
    r_bullet.font.color.rgb = RGBColor(27, 54, 93)
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(40, 40, 40)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(50, 50, 50)
    return p

def create_styled_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        cell = hdr_cells[i]
        set_cell_background(cell, "1B365D") # Navy blue
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header_text)
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    # Data Rows
    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F9FBFD" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            cell = row_cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(cell_value)
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(50, 50, 50)
            
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table

def main():
    # Load fresh copy of template by recreating document or modifying Project Plan.docx
    doc_path = "Project Plan.docx"
    doc = docx.Document(doc_path)
    
    # Clean any previously appended paragraphs and tables after Table 0
    # Keep only paragraph 0..3 and Table 0
    for p in doc.paragraphs[4:]:
        p._p.getparent().remove(p._p)
    for t in doc.tables[1:]:
        t._tbl.getparent().remove(t._tbl)
        
    # 1. Update Existing Team Table (Ensure 5 members + 1 header = 6 rows)
    table = doc.tables[0]
    set_table_borders(table)
    
    # If table has fewer than 6 rows, add rows
    while len(table.rows) < 6:
        table.add_row()
    
    # Style existing header
    for cell in table.rows[0].cells:
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        p = cell.paragraphs[0]
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            
    team_members = [
        ("Andy Wei Jian CHU", "104401283", 
         "Member 1 (Data Engineering & Robustness Lead - 20%):\n"
         "• Data pipeline engineering (loading DPD-Small, preprocessing, stratified splits, augmentations).\n"
         "• Developing real-world field perturbation benchmark suite (lighting shifts, blur, foliage clutter).\n"
         "• Implementing domain-generalization regularization (CutMix, MixUp, RandAugment).\n"
         "• Drafting Report Section 1 (Problem Formulation & EDA) & Section 3.3 (Robustness Evaluation).\n"
         "• Presentation video co-speaker (Data & Robustness)."),
         
        ("Fa Zhen NG", "104400141", 
         "Member 2 (Multi-Task & Independent Baseline Lead - 20%):\n"
         "• GitHub repository setup, branching rules, and inviting lecturer (shlee@swinburne.edu.my).\n"
         "• Implementing & training Baseline 1 (Independent CNNs for crop and disease classification).\n"
         "• Implementing & training Baseline 2 (Multi-Task CNN with shared backbone & dual heads).\n"
         "• Multi-task loss tuning (L_total = α·L_crop + β·L_disease) and Grad-CAM interpretability analysis.\n"
         "• Drafting Report Section 2.1 (Baseline Architectures) & Presentation co-speaker (Baselines)."),
         
        ("Member 3", "[Student ID 3]", 
         "Member 3 (Advanced Deep Learning & Proposed Zero-Shot Model Lead - 20%):\n"
         "• Research & architecture formulation for the Proposed New Approach (Vision-Language CLIP / Contrastive Disentanglement).\n"
         "• Designing semantic attribute space & cross-modal contrastive alignment loss.\n"
         "• Model training, hyperparameter optimization, and evaluating zero-shot transfer on 10 unseen classes.\n"
         "• Drafting Report Section 2.2 (Proposed Methodology) & Section 3.2 (Results & Discussion Lead).\n"
         "• Report academic proofreading & Presentation co-speaker (Proposed Approach)."),
         
        ("Member 4", "[Student ID 4]", 
         "Member 4 (Evaluation Harness, Benchmarking & Error Analysis Lead - 20%):\n"
         "• Automated evaluation harness development (Top-1 & Top-5 accuracy calculation for seen and unseen sets).\n"
         "• Implementing Harmonic Mean (H) metric calculator and high-resolution confusion matrix generator.\n"
         "• Comparative benchmarking across all models, cross-model ablation studies, and qualitative failure clustering.\n"
         "• Drafting Report Section 3.1 (Evaluation Framework) & compiling publication figures/tables.\n"
         "• Cover sheet, checklist, 'Who did what' compliance & Presentation co-speaker (Evaluation)."),

        ("Member 5", "[Student ID 5]", 
         "Member 5 (UI Engineering & Image Retrieval Specialist - 20%):\n"
         "• Designing and implementing interactive Graphical User Interface (Streamlit / Gradio Web App).\n"
         "• Developing deep feature extraction pipeline & vector similarity search (FAISS / cosine distance) over training set.\n"
         "• Building UI gallery component to retrieve and render top matching training samples alongside predictions.\n"
         "• End-to-end UI latency optimization (<2s), error handling, and app packaging.\n"
         "• Drafting Report Section 2.3 (UI Systems), Slide deck lead & Video recording/editing lead (max 8 mins).")
    ]
    
    col_widths = [1.6, 1.1, 4.0]
    for r_idx, (name, sid, contrib) in enumerate(team_members):
        row = table.rows[r_idx + 1]
        row.cells[0].text = name
        row.cells[1].text = sid
        row.cells[2].text = contrib
        bg_color = "F9FBFD" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell in enumerate(row.cells):
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(50, 50, 50)
                if c_idx == 0:
                    run.font.bold = True
                    
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)

    # 2. Project Overview & Problem Definition
    add_heading_1(doc, "1. Project Overview & Problem Formulation")
    add_body_paragraph(
        doc,
        "This project investigates Out-of-Distribution (OOD) and Unseen Plant Disease Identification in the COS30082 "
        "Applied Machine Learning curriculum. In dynamic agricultural environments, automated visual diagnostic systems "
        "frequently encounter unseen disease categories, emerging pathogens, and severe environmental variations (lighting, "
        "complex background foliage, and varied camera viewpoints) not present during model training."
    )
    add_body_paragraph(
        doc,
        "The central research challenge is to achieve a balance between discriminative ability and generalization capability. "
        "A successful model must capture disease-specific features on seen classes while learning transferable representations "
        "capable of zero-shot recognition on unseen categories. The project deliverables include an operational interactive graphical "
        "user interface (GUI), an 8-page academic report, and an 8-minute presentation video, delivered by a 5-member team with "
        "an exact 20% balanced workload distribution."
    )
    
    add_heading_2(doc, "1.1 Dataset Specification (DPD-Small)")
    add_body_paragraph(
        doc,
        "All experiments and benchmarking are conducted using the curated Deep-Plant-Disease Small (DPD-Small) dataset, consisting of 11,448 real-world images:",
        bold_prefix="Benchmark Dataset: "
    )
    
    dataset_headers = ["Split Name", "Image Count", "Metadata File", "Description"]
    dataset_data = [
        ["Train Set", "8,000", "DPD-Small train.csv", "Training images from seen crop-disease categories (323 composition classes across 55 crops and 175 diseases)."],
        ["Test Seen Set", "2,000", "DPD-Small test seen.csv", "Evaluation images belonging to seen crop-disease composition classes to measure discriminative accuracy."],
        ["Test Unseen Set", "1,448", "DPD-Small test unseen.csv", "Evaluation images belonging to 10 unseen crop-disease composition categories (Zero-Shot / OOD)."],
        ["Total Dataset", "11,448", "Combined Metadata", "7 metadata columns: filename, crop label, disease label, composition label, crop name, disease name, composition name."]
    ]
    create_styled_table(doc, dataset_headers, dataset_data, col_widths=[1.3, 1.0, 1.8, 2.6])

    # 3. Requirements Specification
    add_heading_1(doc, "2. Project Requirements Specification")
    
    add_heading_2(doc, "2.1 Basic Requirements (Weight: 12% + Administrative)")
    add_bullet_point(doc, "Train two distinct CNN backbones: Model A classifies 55 crop classes, while Model B classifies 175 disease categories. Both utilize standard data preprocessing, transfer learning, and augmentations.", bold_prefix="Baseline 1 (Independent CNNs): ")
    add_bullet_point(doc, "Train a single shared deep learning backbone with multi-task learning heads simultaneously predicting crop and disease classes with a balanced composite loss function.", bold_prefix="Baseline 2 (Multi-Task CNN): ")
    add_bullet_point(doc, "Provide Top-1 and Top-5 accuracy on both seen and unseen test sets.", bold_prefix="Accuracy Metrics: ")
    add_bullet_point(doc, "Compute the Harmonic Mean between seen and unseen accuracies: H = (2 × Seen × Unseen) / (Seen + Unseen).", bold_prefix="Harmonic Mean (H): ")
    add_bullet_point(doc, "Generate and analyze high-resolution confusion matrix for seen test predictions.", bold_prefix="Confusion Matrix: ")
    add_bullet_point(doc, "Grant read-only access to shlee@swinburne.edu.my within 1 week of team formation.", bold_prefix="Git Repository: ")
    add_bullet_point(doc, "Deliver an 8-page academic report (Methodology, Results & Discussion) and an 8-minute presentation video link by 11:59 pm, 27/11/2026 (Friday Week 12).", bold_prefix="Deliverables: ")

    add_heading_2(doc, "2.2 High-Level Requirements (Weight: 13% + 5% UI)")
    add_bullet_point(doc, "Propose, implement, and validate an advanced deep learning method specifically engineered to surpass conventional baselines in unseen class generalization (U) and environmental robustness.", bold_prefix="Proposed Novel Approach (13%): ")
    add_bullet_point(doc, "Formulate knowledge transfer mechanisms using Vision-Language Zero-Shot models (e.g., CLIP / BioCLIP adapters), Contrastive Representation Learning, or Compositional Disentanglement.", bold_prefix="Zero-Shot Formulation: ")
    add_bullet_point(doc, "Evaluate resilience against real-world field perturbations (variations in lighting conditions, complex background foliage, camera viewpoints, and image noise).", bold_prefix="Robustness Evaluation: ")
    add_bullet_point(doc, "Develop an interactive, user-friendly graphical application (Streamlit or Gradio). The UI takes a plant disease photo as input, outputs the predicted crop-disease category with confidence scores, and retrieves/displays the most visually/semantically similar matching training images.", bold_prefix="User Interface & Retrieval (5%): ")

    # 4. Timeline and Sprints
    add_heading_1(doc, "3. Project Timeline & Two-Sprint Agile Roadmap (Weeks 5 – 12)")
    add_body_paragraph(
        doc,
        "The project spans an 8-week timeline starting at Week 5 and concluding at Week 12. Active technical development "
        "is organized into two structured sprints (Weeks 5 to 10), while the final two weeks (Weeks 11 & 12) are dedicated to "
        "reviewing back, cross-verifying experimental results, authoring the 8-page academic report, and recording the 8-minute presentation video."
    )
    
    add_heading_2(doc, "3.1 Sprint 1: Data Foundations & Baseline Models (Weeks 5 – 7)")
    sprint1_headers = ["Week", "Sprint Focus", "5-Member Technical Activities & Tasks", "Milestone / Deliverable"]
    sprint1_data = [
        ["Week 5", "Setup, EDA & Pipeline Setup", 
         "• M1: Validate DPD-Small dataset, inspect class distributions, build PyTorch Dataset & DataLoader.\n"
         "• M2: Initialize Git repo, branch rules, invite lecturer (shlee@swinburne.edu.my), prepare baseline scaffold.\n"
         "• M3: Literature review on zero-shot plant disease recognition; define text prompt / semantic attribute space.\n"
         "• M4: Build automated evaluation harness (Top-1/5, Harmonic Mean formula, validation scripts).\n"
         "• M5: Wireframe UI layouts (Streamlit/Gradio) & write baseline model inference wrapper script.",
         "Operational data loaders, Git repo shared, and evaluation harness ready."],
        ["Week 6", "Baseline Models Development", 
         "• M1: Implement augmentation pipelines (CutMix, MixUp, RandAugment) and create stratified validation splits.\n"
         "• M2: Train Baseline 1 (Independent CNNs) and Baseline 2 (Multi-Task CNN with shared backbone & dual heads).\n"
         "• M3: Design Proposed Approach architecture (Vision-Language / Contrastive Disentanglement) & zero-shot mapping.\n"
         "• M4: Implement seen test confusion matrix generator & automated metric logging (TensorBoard/CSV).\n"
         "• M5: Build initial UI prototype (image uploader & result containers) linked to baseline checkpoint.",
         "Trained checkpoints for Baseline 1 & 2; initial working UI prototype."],
        ["Week 7", "Baseline Benchmarking & Sprint Review", 
         "• M1: Construct synthetic field perturbation test suite (lighting shifts, blur, foliage background clutter).\n"
         "• M2: Finalize Baseline 1 & Baseline 2 training runs, log loss curves, extract Grad-CAM attention maps.\n"
         "• M3: Implement pre-training & fine-tuning loop for Proposed Model image-text alignment.\n"
         "• M4: Run full benchmark on Baselines (seen/unseen Top-1/5, Harmonic Mean H, confusion matrix).\n"
         "• M5: Implement feature extraction script to compute deep embeddings for all 8,000 training images.\n"
         "• ALL: Conduct Sprint 1 review, analyze baseline generalization gap, draft Report Section 2.1.",
         "Complete baseline benchmark table & draft Baseline report section."]
    ]
    create_styled_table(doc, sprint1_headers, sprint1_data, col_widths=[0.8, 1.3, 3.2, 1.4])

    add_heading_2(doc, "3.2 Sprint 2: Proposed Approach & User Interface (Weeks 8 – 10)")
    sprint2_headers = ["Week", "Sprint Focus", "5-Member Technical Activities & Tasks", "Milestone / Deliverable"]
    sprint2_data = [
        ["Week 8", "Proposed Model Design & Retrieval Index", 
         "• M1: Run field robustness perturbation stress tests on Baseline 1 & 2; quantify vulnerability to noise.\n"
         "• M2: Hyperparameter tuning and loss weight exploration (α, β) for Multi-Task baseline.\n"
         "• M3: Implement complete Proposed Approach (CLIP adapter or Contrastive Disentangled Network).\n"
         "• M4: Expand evaluation harness to compute zero-shot accuracy over the 10 unseen classes (U).\n"
         "• M5: Build vector similarity search index (FAISS / cosine distance) over training image embeddings.",
         "Proposed model functional prototype & vector retrieval index ready."],
        ["Week 9", "Model Training & UI Integration", 
         "• M1: Apply domain-generalization augmentations (RandAugment, background jitter) to Proposed Model.\n"
         "• M2: Analyze baseline model complexity (FLOPs, parameter count, memory footprint, throughput).\n"
         "• M3: Train and fine-tune Proposed Model; optimize contrastive temperature and projection heads.\n"
         "• M4: Evaluate Proposed Model on seen & unseen sets; compute Harmonic Mean H & compare to baselines.\n"
         "• M5: Integrate Proposed Model and image retrieval engine into UI; display top matching training photos.",
         "Converged Proposed Model & functional UI displaying matching training samples."],
        ["Week 10", "Comparative Evaluation & UI Polish", 
         "• M1: Benchmark Proposed Model against field perturbations; verify robustness superiority over baselines.\n"
         "• M2: Consolidate cross-model comparative data and compile comparative architectural analysis.\n"
         "• M3: Document ablation studies on loss components and analyze zero-shot knowledge transfer.\n"
         "• M4: Produce final benchmark tables (Top-1, Top-5, H) and perform qualitative error clustering.\n"
         "• M5: Optimize UI response latency (<2s), implement error handling, and package single-command launch.\n"
         "• ALL: Sprint 2 Milestone Review: verify all experimental numbers and finalize technical assets.",
         "Final comparative benchmark table & production-ready UI."]
    ]
    create_styled_table(doc, sprint2_headers, sprint2_data, col_widths=[0.8, 1.3, 3.2, 1.4])

    add_heading_2(doc, "3.3 Review Back & Finalization Phase (Weeks 11 – 12)")
    review_headers = ["Week", "Sprint Focus", "5-Member Technical Activities & Tasks", "Milestone / Deliverable"]
    review_data = [
        ["Week 11", "Error Analysis & Report Drafting", 
         "• M1: Draft Report Section 1 (Problem Formulation & EDA) & Section 3.3 (Robustness Evaluation).\n"
         "• M2: Draft Report Section 2.1 (Baseline 1 & 2 Methodology, Loss Functions, Training Scheme).\n"
         "• M3: Draft Report Section 2.2 (Proposed Methodology) & Section 3.2 (Results & Discussion Lead).\n"
         "• M4: Draft Report Section 3.1 (Evaluation Setup) & compile publication-quality figures/tables.\n"
         "• M5: Draft Report Section 2.3 (UI & Retrieval Systems) & prepare presentation slide deck structure.\n"
         "• ALL: Review Back: cross-verify results, perform peer-editing to enforce strict 8-page limit.",
         "Complete draft of 8-page report & presentation slides."],
        ["Week 12", "Video Recording & Canvas Submission", 
         "• M1: Clean repository code, standardize docstrings, test clean-slate reproducibility script.\n"
         "• M2: Audit Git commit history and confirm lecturer read-only access (shlee@swinburne.edu.my).\n"
         "• M3: Perform rigorous academic proofreading of report (equations, citations, justifications).\n"
         "• M4: Complete cover sheet, checklist, and 'Who did what' declaration (exact 20% balance per member).\n"
         "• M5: Lead recording and video editing of 8-minute presentation video (all 5 members present); embed URL in report.\n"
         "• ALL: Final joint submission review meeting; submit report to Canvas before 11:59 pm, 27/11/2026.",
         "Final Canvas submission & published 8-minute video presentation."]
    ]
    create_styled_table(doc, review_headers, review_data, col_widths=[0.8, 1.3, 3.2, 1.4])

    # 5. Balanced Task Distribution Matrix
    add_heading_1(doc, "4. Balanced 5-Member Workload Distribution Matrix (20% Each)")
    add_body_paragraph(
        doc,
        "To guarantee an equitable 20% distribution of effort among all five group members, responsibilities are partitioned "
        "into distinct primary ownership and collaborative secondary roles across data engineering, baseline architectures, "
        "advanced zero-shot research, evaluation harness, and UI/retrieval systems."
    )
    
    alloc_headers = ["Group Member & Role", "Sprint 1 Focus (Weeks 5–7)", "Sprint 2 Focus (Weeks 8–10)", "Review Back (Weeks 11–12)", "Contribution"]
    alloc_data = [
        ["Andy Wei Jian CHU\n(M1: Data & Robustness Lead)",
         "• DPD-Small validation & EDA\n• PyTorch Dataset/DataLoader\n• CutMix/MixUp/RandAugment",
         "• Field perturbation suite\n• Robustness stress testing\n• Domain shift evaluation",
         "• Report Sec 1 & Sec 3.3\n• Code cleanup & README\n• Presentation co-speaker",
         "20%\n(Balanced)"],
         
        ["Fa Zhen NG\n(M2: Multi-Task & Baseline Lead)",
         "• GitHub repo setup & rules\n• Baseline 1 (Independent CNNs)\n• Baseline 2 (Multi-Task CNN)",
         "• Multi-task loss tuning (α, β)\n• Grad-CAM visual maps\n• Model complexity & FLOPs",
         "• Report Sec 2.1 (Baselines)\n• Git access compliance\n• Presentation co-speaker",
         "20%\n(Balanced)"],
         
        ["Member 3\n(M3: Advanced DL Researcher)",
         "• Zero-shot literature review\n• Semantic attribute mapping\n• Proposed Approach design",
         "• Proposed Model implementation\n• Contrastive loss & training\n• Unseen class zero-shot tuning",
         "• Report Sec 2.2 & 3.2 Lead\n• Academic proofreading\n• Presentation co-speaker",
         "20%\n(Balanced)"],
         
        ["Member 4\n(M4: Evaluation & Metrics Lead)",
         "• Evaluation harness (Top-1/5)\n• Harmonic Mean (H) calculator\n• Seen test confusion matrix",
         "• Unseen zero-shot benchmark\n• Comparative metric tables\n• Failure cluster analysis",
         "• Report Sec 3.1 & Figures\n• 'Who did what' form\n• Presentation co-speaker",
         "20%\n(Balanced)"],

        ["Member 5\n(M5: UI & Retrieval Specialist)",
         "• UI layout wireframing\n• Baseline inference wrapper\n• Feature embedding extraction",
         "• FAISS/cosine vector search\n• UI retrieval gallery\n• Latency tuning (<2s) & UX",
         "• Report Sec 2.3 (UI Systems)\n• Slide deck lead\n• Video editing lead (≤8 mins)",
         "20%\n(Balanced)"]
    ]
    create_styled_table(doc, alloc_headers, alloc_data, col_widths=[1.5, 1.4, 1.4, 1.4, 0.9])

    # 6. Evaluation Protocol & Quality Control
    add_heading_1(doc, "5. Evaluation Protocol & Quality Assurance Plan")
    add_body_paragraph(
        doc,
        "System evaluation strictly adheres to the course specifications outlined in the project brief:"
    )
    add_bullet_point(doc, "Measures whether the ground truth label is the top prediction (Top-1) or within the top 5 predicted classes (Top-5). Evaluated separately on the 2,000 seen test images and 1,448 unseen test images.", bold_prefix="Top-1 and Top-5 Accuracy: ")
    add_bullet_point(doc, "Calculated as H = (2 × Seen Accuracy × Unseen Accuracy) / (Seen Accuracy + Unseen Accuracy). Penalizes models that perform well on seen categories but fail to generalize to unseen categories.", bold_prefix="Harmonic Mean (H): ")
    add_bullet_point(doc, "Detailed breakdown of predictions on seen test images to identify common confusion clusters between visually similar diseases and crops.", bold_prefix="Confusion Matrix: ")
    add_bullet_point(doc, "End-to-end inference including image preprocessing, prediction, and vector retrieval of matching training images must execute in under 2 seconds on standard GPU/CPU hardware.", bold_prefix="UI Response Latency: ")
    add_bullet_point(doc, "Strict adherence to constraints: Report ≤ 8 pages, Video ≤ 8 minutes, Git read-only access verified, submission before 11:59 pm on 27/11/2026.", bold_prefix="Submission Verification: ")

    # Save document
    try:
        doc.save(doc_path)
        print(f"Successfully saved updated 5-member project plan to {doc_path}")
    except PermissionError:
        fallback_path = "Project Plan (5 Members).docx"
        doc.save(fallback_path)
        print(f"Project Plan.docx is currently locked (open in Word). Successfully saved to {fallback_path}")

if __name__ == "__main__":
    main()
